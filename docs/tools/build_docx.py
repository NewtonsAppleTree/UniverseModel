"""Build the design document from the maintained Markdown sources."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
PROFILES = {
    name: (tuple(ROOT / source for source in profile["sources"]),
           profile["filename"], profile["title"], profile["subtitle"])
    for name, profile in json.loads(
        (ROOT / "docs/tools/documents.json").read_text(encoding="utf-8")
    ).items()
}
FONT = "Microsoft YaHei"
BLACK = RGBColor(0, 0, 0)


def set_font(style, size: float, *, bold: bool = False) -> None:
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = False
    style.font.underline = False
    style.font.color.rgb = BLACK
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attr}"), FONT)


def inline(paragraph, text: str) -> None:
    """Handle the inline code and emphasis used in the source documents."""
    for part in re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text):
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run._element.get_or_add_rPr().get_or_add_rFonts().set(
                qn("w:eastAsia"), FONT
            )
        elif part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        else:
            paragraph.add_run(part)


def configure(document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.9)
    section.bottom_margin = Cm(1.9)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)
    section.footer_distance = Cm(0.9)
    for grid in list(section._sectPr.findall(qn("w:docGrid"))):
        section._sectPr.remove(grid)

    styles = document.styles
    # Strip decorative paragraph borders inherited from the bundled template.
    for border in list(styles.element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
    for border in list(document.element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
    set_font(styles["Normal"], 10.5)
    styles["Normal"].paragraph_format.line_spacing = Pt(16)
    styles["Normal"].paragraph_format.space_after = Pt(4)
    styles["Normal"].paragraph_format.widow_control = True
    snap = OxmlElement("w:snapToGrid")
    snap.set(qn("w:val"), "0")
    styles["Normal"].element.get_or_add_pPr().append(snap)

    for name, size in (("Title", 21), ("Subtitle", 10.5),
                       ("Heading 1", 15), ("Heading 2", 12)):
        set_font(styles[name], size, bold=name != "Subtitle")
        fmt = styles[name].paragraph_format
        fmt.keep_with_next = True
        fmt.space_before = Pt(12 if name.startswith("Heading") else 0)
        fmt.space_after = Pt(7)
        fmt.line_spacing = Pt(size * 1.35)
    styles["Title"].paragraph_format.space_after = Pt(6)
    styles["Subtitle"].paragraph_format.space_after = Pt(12)

    for name in ("List Bullet", "List Number"):
        set_font(styles[name], 10.5)
        styles[name].paragraph_format.line_spacing = Pt(16)
        styles[name].paragraph_format.space_after = Pt(4)

    table_style = styles.add_style("Table Body", WD_STYLE_TYPE.PARAGRAPH)
    table_style.base_style = styles["Normal"]
    set_font(table_style, 9.5)
    table_style.paragraph_format.line_spacing = Pt(14)
    table_style.paragraph_format.space_before = Pt(0)
    table_style.paragraph_format.space_after = Pt(0)

    code_style = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    code_style.base_style = styles["Normal"]
    set_font(code_style, 9)
    code_style.font.name = "Consolas"
    code_style.paragraph_format.line_spacing = 1.05
    code_style.paragraph_format.space_after = Pt(0)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "18")
    props.append(size)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    field.append(run)
    footer._p.append(field)


def add_table(document, lines: list[str], *, reference: bool = False) -> None:
    rows = []
    for line in lines:
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            continue
        rows.append(cells)
    columns = len(rows[0])
    if any(len(row) != columns for row in rows):
        raise ValueError("Inconsistent Markdown table column count")

    widths = {2: (4.0, 13.0), 3: (3.5, 6.3, 7.2),
              4: (3.5, 3.0, 3.5, 7.0)}.get(columns)
    if widths is None:
        raise ValueError(f"Unsupported table width: {columns}")
    if rows[0][0] == "阶段":
        widths = (1.3, 6.2, 9.5)
    elif not reference and rows[0][0] == "对象":
        widths = (4.3, 5.7, 7.0)
    elif not reference and rows[0][0] == "关系":
        widths = (4.2, 5.6, 7.2)
    elif reference and columns == 3:
        widths = (4.6, 6.4, 6.0)

    if document.paragraphs and document.paragraphs[-1].text:
        document.paragraphs[-1].paragraph_format.keep_with_next = True

    table = document.add_table(rows=0, cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = Cm(width)

    properties = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{side}")
        for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            border.set(qn(f"w:{key}"), value)
        borders.append(border)
    properties.append(borders)
    margins = OxmlElement("w:tblCellMar")
    for side, value in (("top", "50"), ("bottom", "50"),
                        ("left", "100"), ("right", "100")):
        item = OxmlElement(f"w:{side}")
        item.set(qn("w:w"), value)
        item.set(qn("w:type"), "dxa")
        margins.append(item)
    properties.append(margins)

    for index, values in enumerate(rows):
        row = table.add_row()
        row_properties = row._tr.get_or_add_trPr()
        row_properties.append(OxmlElement("w:cantSplit"))
        if index == 0:
            row_properties.append(OxmlElement("w:tblHeader"))
        for column_index, (cell, value, width) in enumerate(zip(row.cells, values, widths)):
            cell.width = Cm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.style = "Table Body"
            inline(paragraph, value)
            if rows[0][0] == "阶段" and column_index == 0:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if index == 0:
                paragraph.paragraph_format.keep_with_next = True
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "E8EEF4")
                cell._tc.get_or_add_tcPr().append(shading)
                for run in paragraph.runs:
                    run.bold = True
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(4)
    spacer.add_run().font.size = Pt(4)


def append_markdown(document, content: str, section_number: int,
                    *, reference: bool = False,
                    page_break_titles: tuple[str, ...] = ()) -> None:
    lines = content.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith("# "):
            if not reference:
                paragraph = document.add_paragraph(
                    f"{section_number} {line[2:]}", "Heading 1"
                )
                if section_number > 1:
                    paragraph.paragraph_format.page_break_before = True
        elif line.startswith("## "):
            heading = line[3:]
            if not reference:
                heading = re.sub(r"^(\d+)\s", rf"{section_number}.\1 ", heading)
            paragraph = document.add_paragraph(heading, "Heading 1" if reference else "Heading 2")
            if line[3:] in page_break_titles:
                paragraph.paragraph_format.page_break_before = True
        elif line.startswith("### "):
            document.add_paragraph(line[4:], "Heading 2")
        elif line.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index])
                index += 1
            add_table(document, table_lines, reference=reference)
            continue
        elif line.startswith("```"):
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                document.add_paragraph(lines[index], "Code Block")
                index += 1
        elif re.match(r"^\d+\.\s", line):
            paragraph = document.add_paragraph(style="List Number")
            inline(paragraph, re.sub(r"^\d+\.\s+", "", line))
        elif line.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            inline(paragraph, line[2:])
        else:
            paragraph = document.add_paragraph()
            inline(paragraph, line)
        index += 1


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--document", choices=tuple(PROFILES), default="design")
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    reference = args.document != "design"
    sources, filename, title, subtitle = PROFILES[args.document]
    contents = [path.read_text(encoding="utf-8") for path in sources]
    digest = hashlib.sha256("\n".join(contents).encode("utf-8")).hexdigest()
    if args.output is None:
        args.output = ROOT / "docs" / filename

    document = Document()
    configure(document)
    if reference:
        document.sections[0].page_width = Inches(8.5)
        document.sections[0].page_height = Inches(11)
        set_font(document.styles["Normal"], 11)
        document.styles["Normal"].paragraph_format.line_spacing = Pt(17)
    document.add_paragraph(title, "Title")
    document.add_paragraph(subtitle, "Subtitle")
    for number, content in enumerate(contents, start=1):
        append_markdown(document, content, number, reference=reference,
                        page_break_titles=("11 参考资料",) if args.document == "scaling" else ())

    document.core_properties.title = title
    document.core_properties.subject = subtitle
    document.core_properties.author = "UniverseModel"
    document.core_properties.comments = f"Markdown source SHA256: {digest}"
    args.output.parent.mkdir(parents=True, exist_ok=True)
    document.save(args.output)
    print(args.output)
    print(f"Source SHA256: {digest}")


if __name__ == "__main__":
    main()
